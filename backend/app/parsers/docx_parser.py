from io import BytesIO

from docx import Document

from app.parsers.base_parser import BaseResumeParser 

class DOCXResumeParser(BaseResumeParser):


    async def extract_text(
            self,
            *,
            file_bytes: bytes,
    ) -> str:

        try:
            document = Document(BytesIO(file_bytes))

        except Exception as exc:
            raise ValueError(
                "The uploaded file is not a valid DOCX document."
            ) from exc

        text_parts: list[str] = [""]

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                text_parts.append(text)


        # tables
        for table in document.tables:

            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    cell_text = cell = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)
                if row_text:
                    text_parts.append(
                        " | ".join(row_text)
                    )

        return "\n".join(text_parts)
    